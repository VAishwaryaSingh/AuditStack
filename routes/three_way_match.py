from flask import Blueprint, render_template, request, jsonify
from services.database import (
    get_match_sets, create_match_set, delete_match_set,
    get_match_transactions,
    upsert_match_po, upsert_match_grn, upsert_match_invoice,
)
import csv, io, re

three_way_match_bp = Blueprint("three_way_match", __name__)


# ── HELPERS ────────────────────────────────────────────────────────

def _col(row, *keys):
    lower = {k.strip().lower(): v for k, v in row.items()}
    for k in keys:
        v = lower.get(k.lower(), "")
        if v and str(v).strip():
            return str(v).strip()
    return ""


def _parse_csv(f):
    raw = f.read()
    try:
        text = raw.decode("utf-8-sig")
    except Exception:
        text = raw.decode("latin-1")
    reader = csv.DictReader(io.StringIO(text))
    return list(reader)


def _detect_doc_type(headers):
    """Return 'po'|'grn'|'invoice'|None based on column headers."""
    h = {h.strip().lower() for h in headers if h}
    grn_hits     = h & {'received_by', 'received by', 'receieved_by', 'receivedby', 'location', 'warehouse'}
    invoice_hits = h & {'due_date', 'due date', 'payment_due', 'payment due', 'invoice_date', 'invoice date', 'bill date'}
    po_hits      = h & {'buyer', 'requested_by', 'requested by', 'purchase order', 'po_date', 'po date', 'ordered_by'}
    scores = {'grn': len(grn_hits), 'invoice': len(invoice_hits), 'po': len(po_hits)}
    best = max(scores, key=scores.get)
    return best if scores[best] > 0 else None


_DOC_LABELS = {'po': 'Purchase Orders', 'grn': 'Goods Received Notes', 'invoice': 'Supplier Invoices'}


def _parse_pdf(file_bytes, doc_type):
    """Best-effort PDF text extraction → list of field dicts."""
    try:
        import pdfplumber
        text = ""
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                text += (page.extract_text() or "") + "\n"
    except Exception as e:
        raise ValueError(f"Could not read PDF: {e}")

    row = {}

    # transaction/reference number
    for pat in [
        r'(?:PO|Purchase Order|Order|GRN|Goods Received|Invoice|INV|TXN|Transaction|Ref(?:erence)?)[^\w\n]*[#No\.:\s]+([A-Z0-9][A-Z0-9\-/]{2,})',
        r'\b((?:PO|GRN|INV|TXN)-[\w\-]+)\b',
    ]:
        m = re.search(pat, text, re.IGNORECASE)
        if m:
            row.setdefault('ref', m.group(1).strip())
            row.setdefault('transaction_id', m.group(1).strip())
            break

    # vendor / supplier
    for pat in [r'(?:Vendor|Supplier|To|Bill To|Sold To|From)[:\s]+([A-Za-z0-9][A-Za-z0-9\s,\.&\-]{2,60})']:
        m = re.search(pat, text, re.IGNORECASE)
        if m:
            row['vendor'] = m.group(1).strip()[:100]
            break

    # date
    for pat in [
        r'(?:Date|Order Date|Invoice Date|GRN Date|Dated)[:\s]+(\d{1,2}[\/\-\.]\d{1,2}[\/\-\.]\d{2,4})',
        r'\b(\d{4}-\d{2}-\d{2})\b',
    ]:
        m = re.search(pat, text, re.IGNORECASE)
        if m:
            row['date'] = m.group(1).strip()
            break

    # quantity
    m = re.search(r'(?:Qty|Quantity|Units?)[:\s]+(\d[\d,]*(?:\.\d+)?)', text, re.IGNORECASE)
    if m:
        row['quantity'] = m.group(1).replace(',', '').strip()

    # unit price
    m = re.search(r'(?:Unit Price|Price|Rate)[:\s]+[\$£€]?\s*(\d[\d,]*(?:\.\d+)?)', text, re.IGNORECASE)
    if m:
        row['unit_price'] = m.group(1).replace(',', '').strip()

    # line total
    for pat in [r'(?:Line Total|Total Amount|Amount|Total)[:\s]+[\$£€]?\s*(\d[\d,]*(?:\.\d+)?)']:
        m = re.search(pat, text, re.IGNORECASE)
        if m:
            row['line_total'] = m.group(1).replace(',', '').strip()
            break

    # doc-type-specific fields
    if doc_type == 'grn':
        m = re.search(r'(?:Received By|Receiver|Warehouse|Location)[:\s]+([A-Za-z0-9][A-Za-z0-9\s\-]{1,50})', text, re.IGNORECASE)
        if m:
            row['received_by'] = m.group(1).strip()
    elif doc_type == 'po':
        m = re.search(r'(?:Buyer|Ordered By|Requested By|Prepared By)[:\s]+([A-Za-z][A-Za-z\s\-]{1,50})', text, re.IGNORECASE)
        if m:
            row['buyer'] = m.group(1).strip()
    elif doc_type == 'invoice':
        m = re.search(r'(?:Due Date|Payment Due|Pay By)[:\s]+(\d{1,2}[\/\-\.]\d{1,2}[\/\-\.]\d{2,4})', text, re.IGNORECASE)
        if m:
            row['due_date'] = m.group(1).strip()

    return [row] if (row.get('transaction_id') or row.get('ref')) else []


# ── PAGES ──────────────────────────────────────────────────────────

@three_way_match_bp.route("/three-way-match")
def three_way_match_page():
    return render_template("three_way_match.html", match_sets=get_match_sets())


# ── MATCH SETS ─────────────────────────────────────────────────────

@three_way_match_bp.route("/api/match-sets", methods=["POST"])
def api_create_set():
    data = request.get_json(force=True)
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"error": "Name is required"}), 400
    sid = create_match_set(name)
    return jsonify({"id": sid, "name": name}), 201


@three_way_match_bp.route("/api/match-sets/<int:set_id>", methods=["DELETE"])
def api_delete_set(set_id):
    delete_match_set(set_id)
    return jsonify({"ok": True})


@three_way_match_bp.route("/api/match-sets/<int:set_id>/transactions")
def api_get_transactions(set_id):
    return jsonify(get_match_transactions(set_id))


# ── UPLOAD ─────────────────────────────────────────────────────────

@three_way_match_bp.route("/api/match-sets/<int:set_id>/upload", methods=["POST"])
def api_upload(set_id):
    doc_type = request.form.get("doc_type")
    if doc_type not in ("po", "grn", "invoice"):
        return jsonify({"error": "Invalid doc_type"}), 400

    f = request.files.get("file")
    if not f:
        return jsonify({"error": "No file provided"}), 400

    filename = f.filename.lower()
    rows = []
    is_pdf = False

    try:
        if filename.endswith(".pdf"):
            is_pdf = True
            raw = f.read()
            rows = _parse_pdf(raw, doc_type)
            if not rows:
                return jsonify({
                    "error": "Could not extract recognisable fields from this PDF. "
                             "Check that it contains a reference number and upload again, "
                             "or use a CSV export instead.",
                    "type_mismatch": False,
                    "pdf_parse_failed": True,
                }), 422

        elif filename.endswith((".xlsx", ".xls")):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True)
            ws = wb.active
            hdrs = [str(c.value).strip() if c.value is not None else "" for c in next(ws.iter_rows(1, 1))]
            rows = [
                {hdrs[i]: (str(v).strip() if v is not None else "") for i, v in enumerate(row)}
                for row in ws.iter_rows(min_row=2, values_only=True)
                if any(v is not None for v in row)
            ]
        elif filename.endswith(".csv"):
            rows = _parse_csv(f)
        else:
            return jsonify({"error": "Supported formats: CSV, Excel (.xlsx/.xls), PDF"}), 400

    except Exception as e:
        return jsonify({"error": str(e)}), 400

    # Doc-type validation (skip for PDFs — we already parsed knowing the type)
    if not is_pdf and rows:
        detected = _detect_doc_type(rows[0].keys())
        if detected and detected != doc_type:
            return jsonify({
                "error": (
                    f"This file looks like a {_DOC_LABELS[detected]} file "
                    f"(found column '{_grn_or_inv_hint(detected)}'), "
                    f"but you uploaded it to the {_DOC_LABELS[doc_type]} zone. "
                    f"Please move it to the correct upload box."
                ),
                "type_mismatch": True,
                "detected_type": detected,
            }), 422

    count = 0
    for r in rows:
        txn_id = _col(r, "transaction_id", "txn_id", "Transaction ID", "TXN ID", "id", "ref", "reference", "Reference")
        if not txn_id:
            continue
        vendor  = _col(r, "vendor", "Vendor", "supplier", "Supplier")
        ref     = _col(r, "ref", "reference", "Reference", "po_ref", "grn_ref", "inv_ref",
                       "PO Ref", "GRN Ref", "Invoice Ref")
        date    = _col(r, "date", "Date", "po_date", "grn_date", "inv_date")
        item    = _col(r, "line_item", "Line Item", "description", "Description", "item", "Item")
        qty     = _col(r, "quantity", "qty", "Quantity", "QTY", "Qty")
        total   = _col(r, "line_total", "total", "Line Total", "Total", "amount", "Amount")

        if doc_type == "po":
            price = _col(r, "unit_price", "price", "Unit Price", "Price", "Rate", "rate")
            buyer = _col(r, "buyer", "Buyer", "requested_by", "Requested By")
            upsert_match_po(set_id, txn_id, vendor, ref, date, item, qty, price, buyer, total)
        elif doc_type == "grn":
            rec = _col(r, "received_by", "Received By", "received by", "location", "Location", "warehouse", "Warehouse")
            upsert_match_grn(set_id, txn_id, vendor, ref, date, item, qty, rec, total)
        else:
            price    = _col(r, "unit_price", "price", "Unit Price", "Price", "Rate", "rate")
            due_date = _col(r, "due_date", "Due Date", "due date", "payment_due", "Payment Due")
            upsert_match_invoice(set_id, txn_id, vendor, ref, date, item, qty, price, due_date, total)
        count += 1

    return jsonify({"ok": True, "count": count, "is_pdf": is_pdf})


def _grn_or_inv_hint(detected):
    if detected == "grn":
        return "received_by"
    if detected == "invoice":
        return "due_date"
    return "buyer"
