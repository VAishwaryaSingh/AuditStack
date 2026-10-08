from flask import Blueprint, request, jsonify, send_file
import io
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from services.database import (
    create_session, get_sessions, get_session, delete_session,
    upsert_rows, update_explanation, update_threshold
)
from services.parser import parse_upload
from services.calculator import calculate_variances

api_bp = Blueprint("api", __name__)

# Holds parsed DataFrames until both PY and CY are uploaded for a session
PENDING_UPLOADS = {}

@api_bp.route("/api/sessions", methods=["GET"])
def list_sessions():
    return jsonify(get_sessions())

@api_bp.route("/api/sessions", methods=["POST"])
def new_session():
    data = request.get_json() or {}
    name = data.get("name", "Untitled Session").strip() or "Untitled Session"
    threshold = float(data.get("materiality_threshold", 50000.0))
    pct = float(data.get("materiality_pct_threshold", 5.0))
    sid = create_session(name, threshold, pct)
    return jsonify({"id": sid, "name": name}), 201

@api_bp.route("/api/sessions/<int:session_id>", methods=["GET"])
def get_session_data(session_id):
    session = get_session(session_id)
    if not session:
        return jsonify({"error": "Session not found"}), 404
    return jsonify(session)

@api_bp.route("/api/sessions/<int:session_id>", methods=["DELETE"])
def remove_session(session_id):
    delete_session(session_id)
    PENDING_UPLOADS.pop(session_id, None)
    return jsonify({"ok": True})

@api_bp.route("/api/sessions/<int:session_id>/upload", methods=["POST"])
def upload_file(session_id):
    session = get_session(session_id)
    if not session:
        return jsonify({"error": "Session not found"}), 404

    file = request.files.get("file")
    role = request.form.get("role", "").lower()

    if not file or file.filename == "":
        return jsonify({"error": "No file provided"}), 400
    if role not in ("py", "cy"):
        return jsonify({"error": "role must be 'py' or 'cy'"}), 400

    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in ("csv", "xlsx", "xls"):
        return jsonify({"error": "Only CSV and Excel (.xlsx/.xls) files are supported"}), 400

    try:
        file_bytes = file.read()
        df = parse_upload(file_bytes, ext)
    except ValueError as e:
        return jsonify({"error": str(e)}), 422
    except Exception as e:
        return jsonify({"error": f"Failed to parse file: {str(e)}"}), 422

    if session_id not in PENDING_UPLOADS:
        PENDING_UPLOADS[session_id] = {}
    PENDING_UPLOADS[session_id][role] = df

    pending = PENDING_UPLOADS.get(session_id, {})
    if "py" in pending and "cy" in pending:
        rows = calculate_variances(
            pending["py"],
            pending["cy"],
            session["materiality_threshold"],
            session["materiality_pct_threshold"],
        )
        upsert_rows(session_id, rows)
        del PENDING_UPLOADS[session_id]
        return jsonify({"status": "merged", "row_count": len(rows)})

    return jsonify({"status": "uploaded", "role": role, "row_count": len(df)})

@api_bp.route("/api/sessions/<int:session_id>/upload-status", methods=["GET"])
def upload_status(session_id):
    pending = PENDING_UPLOADS.get(session_id, {})
    return jsonify({
        "py_uploaded": "py" in pending,
        "cy_uploaded": "cy" in pending,
    })

@api_bp.route("/api/sessions/<int:session_id>/threshold", methods=["PUT"])
def set_threshold(session_id):
    data = request.get_json() or {}
    threshold = float(data.get("materiality_threshold", 50000.0))
    pct = float(data.get("materiality_pct_threshold", 5.0))
    update_threshold(session_id, threshold, pct)
    session = get_session(session_id)
    return jsonify(session)

@api_bp.route("/api/rows/<int:row_id>/explanation", methods=["PUT"])
def save_explanation(row_id):
    data = request.get_json() or {}
    text = data.get("explanation", "")
    update_explanation(row_id, text)
    return jsonify({"ok": True})

@api_bp.route("/api/sessions/<int:session_id>/export", methods=["GET"])
def export_excel(session_id):
    session = get_session(session_id)
    if not session:
        return jsonify({"error": "Session not found"}), 404

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Risk Assessment"

    header_fill = PatternFill("solid", fgColor="1E3A6E")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    amber_fill = PatternFill("solid", fgColor="FEF3C7")
    red_fill = PatternFill("solid", fgColor="FEE2E2")
    flag_font = Font(color="92400E", bold=True)
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = [
        "Account Code", "Account Name", "PY Balance ($)", "CY Balance ($)",
        "Variance ($)", "Variance (%)", "CTT Status", "Materiality / CTT", "Client Explanation"
    ]
    col_widths = [16, 35, 16, 16, 16, 14, 22, 18, 45]

    for i, (h, w) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 22

    num_fmt = '#,##0.00'
    pct_fmt = '0.00'

    threshold = session["materiality_threshold"]

    for r_idx, row in enumerate(session["rows"], start=2):
        variance = row.get("variance_amount") or 0
        is_flagged = row.get("flagged", 0)
        is_severe = abs(variance) >= threshold * 2

        mat_label = f"${session['materiality_threshold']:,.0f} | {session['materiality_pct_threshold']}%"
        values = [
            row.get("account_code", ""),
            row.get("account_name", ""),
            row.get("py_balance", 0),
            row.get("cy_balance", 0),
            variance,
            row.get("variance_pct"),
            row.get("status", ""),
            mat_label,
            row.get("client_explanation", ""),
        ]
        for c_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=(c_idx == 9))
            if is_severe:
                cell.fill = red_fill
            elif is_flagged:
                cell.fill = amber_fill
            if c_idx in (3, 4, 5):
                cell.number_format = num_fmt
            elif c_idx == 6:
                cell.number_format = pct_fmt
            if is_flagged and c_idx == 7 and not row.get("client_explanation"):
                cell.font = flag_font

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{len(session['rows']) + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    safe_name = "".join(c if c.isalnum() or c in " -_" else "_" for c in session["name"])
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"AuditStack_{safe_name}_Risk_Assessment.xlsx",
    )
